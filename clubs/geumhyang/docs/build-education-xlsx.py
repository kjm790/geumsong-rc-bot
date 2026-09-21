# -*- coding: utf-8 -*-
"""조사 결과(education-videos.md → videos.json) 를 공유드라이브용 엑셀(탭 4개)로 만든다.
링크·제목은 전부 파일에서 읽어 그대로 쓴다(손으로 옮겨 적지 않음)."""
import json, re, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SCR = os.path.dirname(os.path.abspath(__file__))
MD = r"C:\Users\kjm79\ai-samujang-bot\clubs\geumhyang\docs\education-videos.md"
OUT = sys.argv[1]
videos = json.load(open(os.path.join(SCR, "edu", "videos.json"), encoding="utf-8"))
by_id = {v["url"][-11:]: v for v in videos}

# ── 교육과정 초안: (대상, 주제, 주 영상 id, 보조 영상 id들, 정답기준 초안)
CHECK = "시청 후 담당자 작성"
CURR = [
    ("전 임원", "로타리 기초 — 로타리란 무엇인가", "wBzzWmwZUGw", [], "로타리재단 (클럽은 봉사 실행, 국제로타리는 지원·조정, 재단이 자금 지원)"),
    ("전 임원", "네 가지 표준과 로타리안 행동강령", "d5IbX1Zt7gs", ["8TwbigF9gEo"], "①진실한가 ②모두에게 공평한가 ③선의와 우정을 더하게 하는가 ④모두에게 유익한가"),
    ("전 임원", "아호 문화 — 왜 아호로 부르는가", "tr0nuRCAXoU", [], "초안: 직함·서열이 아니라 서로를 동등하게 존중하며 부르기 위함 — " + CHECK),
    ("전 임원", "클럽 임원의 역할 한눈에 보기", "WZxSWuFSCL4", [], "초안: 클럽 운영 자금과 봉사·기부 목적 자금을 섞지 않고 따로 관리하는 것 — " + CHECK),
    ("전 임원", "클럽의 조직 — 상임위원회와 이사회", "EKC1xJEfZoE", ["s2CSlVENt7w"], "청소년(신세대) 위원회"),
    ("전 임원", "표준 정관과 클럽 세칙", "dmSYnJ4w0kE", [], "클럽 세칙 (표준 클럽 정관은 국제로타리가 정한 것이라 클럽이 고칠 수 없음)"),
    ("전 임원", "정기모임 식순과 의전", "EdS9p98yVm8", ["Nn_mJNiIfl4", "TYZuUARlgS4"], "로타리 강령(현 '로타리의 목적') 낭독 → 네 가지 표준 제창"),
    ("전 임원", "로타리재단과 기부 표창(PHF·RFSM)", "CFCe_WTUzTI", ["11Q94Yp9IV8"], "미화 1,000달러"),
    ("전 임원", "신생클럽 창립 절차", "NHbPdXhRbhI", [], "초안: 최소 20명(현행 규정은 지구 사무국 확인), 권장 인원은 " + CHECK),
    ("회장·차기회장·부회장", "직책 연수 — 회장 분과 (3700지구 2026-27)", "eru1GCpviok", ["0i9QjhvOnok"], "3대 역할: 대외 대표·리더·이사회 의장 / 부회장 배치 이유는 " + CHECK),
    ("총무이사", "직책 연수 — 총무 분과 (3700지구 2026-27)", "vZ0i9AD16d4", [], "5대: 클럽관리·멤버십·공공이미지·봉사프로젝트·로타리재단 / 참석 이유는 " + CHECK),
    ("재무이사", "직책 연수 — 재무 분과 (3700지구 2026-27)", "Xs9I1YkqByE", [], "4대 원칙: 투명성·책임성·정확성·윤리성 / 예산 승인 주체는 " + CHECK),
    ("사찰이사", "직책 연수 — 사찰의 임무와 의전", "yaW-oUBpC30", ["TYZuUARlgS4", "dNzkg72K1wU"], CHECK + " (자막 없는 2020년 자료 — 금송RC 사찰이사 실습 병행 권장)"),
    ("클럽관리위원장", "위원회 연수 — 클럽관리", "08YoxbdSYfM", ["s2CSlVENt7w"], CHECK),
    ("멤버십위원장", "위원회 연수 — 멤버십(회원증강)", "a8KHO4lHflw", ["fCVLfN6FtbA", "NHbPdXhRbhI"], "영입 대상 다양화 · 기존 회원 유지 · 클럽 매력도 제고"),
    ("공공이미지위원장", "위원회 연수 — 공공이미지", "q7TiqYmr8Qc", ["lOOrok0fBOg"], CHECK),
    ("로타리재단위원장", "위원회 연수 — 로타리재단", "fJMVm6o5Zpw", ["opZ0eMPfhM0", "CFCe_WTUzTI"], "초안: 3년 뒤 일부가 지구지정기금(DDF)으로 돌아옴 — 정확한 비율은 " + CHECK),
    ("봉사프로젝트위원장", "위원회 연수 — 봉사프로젝트와 7대 초점분야", "rylMQDsgQqM", ["vu9MrQkjsFk"], CHECK),
    ("DEI위원장", "위원회 연수 — DEI(다양성·형평성·포용)", "AjZCcYl6qWQ", ["f7pYyxZgKAY"], CHECK + " · 한국어 영상이 없어 영어 연설(5분)입니다. 러닝센터 강좌 「참여와 소속감 증진하기」·「포용적인 클럽 문화 조성하기」 수강을 함께 안내하세요."),
    ("IT위원장", "위원회 연수 — 클럽 IT위원회와 My Rotary", "DW10P30CO_s", ["F_gUpej7S9s", "84hRN6pXIKQ", "41YIau8a5IY"], CHECK + " · 우리 클럽 IT위원장은 AI 사무장(텔레그램 봇)·공유드라이브 운영 파트너이기도 합니다."),
]
missing = [c[2] for c in CURR if c[2] not in by_id] + [x for c in CURR for x in c[3] if x not in by_id]
if missing:
    raise SystemExit("조사 파일에 없는 영상 id: %s" % missing)

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")
THIN = Side(style="thin", color="D9D9D9")
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(wb, title, headers, rows, widths, link_cols=(), note=None):
    ws = wb.create_sheet(title)
    start = 1
    if note:
        ws.cell(row=1, column=1, value=note).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.row_dimensions[1].height = 48
        start = 2
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=start, column=j, value=h)
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, r in enumerate(rows, start + 1):
        for j, val in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=val)
            c.alignment, c.border = WRAP, Border(bottom=THIN)
            if j in link_cols and isinstance(val, str) and val.startswith("http"):
                c.hyperlink, c.font = val, LINK_FONT
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    ws.auto_filter.ref = "A%d:%s%d" % (start, get_column_letter(len(headers)), start + len(rows))
    return ws


wb = Workbook()
wb.remove(wb.active)

# 1) 교육과정 초안 — 봇 '교육과정' 탭과 같은 열(회차·주제·영상링크·확인질문·정답기준·게시예정일·게시시각) + 대상·길이·보조 영상
rows = []
for n, (who, topic, vid, extra, ans) in enumerate(CURR, 1):
    v = by_id[vid]
    rows.append([n, who, topic, v["url"], v["len"], v["q"], ans, "", "", "\n".join(by_id[x]["url"] for x in extra), v["title"]])
sheet(wb, "교육과정(초안)",
      ["회차", "대상", "주제", "영상링크", "길이", "확인질문", "정답기준(초안·검수 필요)", "게시예정일", "게시시각", "보조 영상", "영상 제목(원문)"],
      rows, [6, 18, 34, 46, 8, 46, 46, 12, 12, 46, 40], link_cols=(4,),
      note="대구금향RC 임원 교육과정 초안 (2026-09-21). 1~9회차는 전 임원 공통, 10회차부터는 직책별(회장단·총무·재무·사찰 + 상임위원장 7)입니다. 정답기준은 조사 요약에서 확인되는 것만 적었고 '시청 후 담당자 작성' 표시는 교육 담당이 영상을 보고 채워야 합니다. "
           "순서·영상은 자유롭게 바꾸세요. 확정되면 봇의 「교육과정」 탭으로 옮겨 회차별로 임원방에 게시합니다.")

# 2) 영상 목록 전체
sheet(wb, "영상목록(72편)",
      ["번호", "분야", "제목", "채널/출처(게시연도)", "언어(자막)", "길이", "링크", "한 줄 요약", "추천 확인질문", "우리 클럽 메모"],
      [[i, v["section"], v["title"], v["channel"], v["lang"], v["len"], v["url"], v["summary"], v["q"], ""] for i, v in enumerate(videos, 1)],
      [6, 22, 44, 24, 18, 8, 46, 60, 46, 24], link_cols=(7,),
      note="모든 링크는 2026-09-21에 실제로 열어 제목·채널·길이를 확인한 것입니다. '제목 기준'이라고 적힌 요약은 자막이 없어 내용을 확인하지 못한 영상입니다 — 교육에 쓰기 전 한 번 시청해 보세요. "
           "2020년 3630지구 자료와 2013년 '권장클럽주회'는 회비 수치·용어가 현행과 다를 수 있습니다(예: 로타리 강령 → 현재 '로타리의 목적').")

# 3) 러닝센터 · 4) 미확인 — md 의 표를 그대로 읽는다
md = open(MD, encoding="utf-8").read().replace("\r\n", "\n")


def table_under(heading, ncols):
    body = md.split("## " + heading, 1)[1].split("\n## ", 1)[0]
    out = []
    for ln in body.split("\n"):
        if not ln.startswith("|") or re.match(r"^\|\s*-{3}", ln):
            continue
        cells = [c.strip() for c in ln.split("|")[1:-1]]
        if len(cells) == ncols:
            out.append(cells)
    return out[1:]  # 제목 줄 제외


courses = table_under("로타리 러닝센터", 4)
sheet(wb, "러닝센터 강좌(23)", ["대상", "강좌(한국어 명칭)", "구성·내용(카탈로그 기준)", "링크(My Rotary 로그인 필요)", "수강 완료 메모"],
      [c + [""] for c in courses], [20, 32, 70, 60, 20], link_cols=(4,),
      note="로타리 학습센터(https://my.rotary.org/learn)는 My Rotary 계정으로 로그인해야 수강할 수 있습니다. 창립(가입 승인) 전 임원의 수강 가능 여부는 지구 사무국에 확인하세요. "
           "사찰 강좌는 학습센터에 없습니다. 강좌명·링크는 국제로타리 한국어 강좌 카탈로그(2026년 7월판)에서 옮긴 것이며 강좌 화면 자체는 로그인 장벽으로 확인하지 못했습니다.")

unv = table_under("미확인/제외", 3)
sheet(wb, "미확인·제외", ["항목", "링크", "사유"], unv, [60, 60, 80],
      note="영상이 실재하는 것은 확인했지만 내용을 검증하지 못했거나 임원 교육용으로 맞지 않아 본 목록에서 뺀 항목입니다. 담당자가 직접 보고 쓸 만하면 영상목록에 추가하세요.")

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("saved", OUT, "| 교육과정", len(rows), "| 영상", len(videos), "| 강좌", len(courses), "| 미확인", len(unv))
