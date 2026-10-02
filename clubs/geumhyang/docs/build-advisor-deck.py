# -*- coding: utf-8 -*-
"""
대구금향RC 「클럽 임원의 역할」 — 신생클럽 어드바이저 교육용 슬라이드 생성기 (python-pptx)
원자료: 3700지구 연수위원회 수경 이명미 「클럽 임원의 역할」(2026-09-16) — 뼈대·문구 요지를 살리고
        금향 사정(13석 이사회, 7대 위원회에 DEI·IT 포함, 초대회장 연임, 창립총회 10/14, AI 사무장)으로 바꿨다.
강사: 신생클럽 어드바이저 심천 김영상(대구금송로타리클럽 4대 회장). 슬라이드마다 강의 메모(노트) 포함.
실행: PYTHONUTF8=1 python build-advisor-deck.py  → 공유드라이브 00_교육자료/ 에 저장
"""
import os, io
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BLUE, GOLD, INK, GRAY, LIGHT, WHITE = RGBColor(0x17, 0x45, 0x8F), RGBColor(0xF7, 0xA8, 0x1B), RGBColor(0x21, 0x25, 0x29), RGBColor(0x6E, 0x73, 0x7D), RGBColor(0xEE, 0xF2, 0xF8), RGBColor(255, 255, 255)
FONT = '맑은 고딕'
OUT_DIR = r'G:\공유 드라이브\대구금향RC\00_교육자료'
OUT = os.path.join(OUT_DIR, '대구금향RC_클럽임원의역할_신생클럽어드바이저교육.pptx')
LOGO = r'C:\Users\kjm79\ai-samujang-bot\videos\로고_대구금향RC.png'
ADVISOR_PHOTO = r'C:\Users\kjm79\ai-samujang-bot\videos\90_신생클럽어드바이저_김영상\photo.png'

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height

def run(p, text, size, bold=False, color=INK):
    r = p.add_run(); r.text = text; f = r.font; f.name = FONT; f.size = Pt(size); f.bold = bold; f.color.rgb = color; return r

def box(sl, x, y, w, h, text='', size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lines=None, spacing=1.15):
    s = sl.shapes.add_shape(shape, x, y, w, h)
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(1.25)
    s.shadow.inherit = False
    tf = s.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.12); tf.margin_top = tf.margin_bottom = Inches(0.06)
    items = lines if lines is not None else ([text] if text else [])
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        if isinstance(it, tuple): run(p, it[0], it[1], it[2] if len(it) > 2 else bold, it[3] if len(it) > 3 else color)
        else: run(p, it, size, bold, color)
    return s

def frame(sl, title, kicker=None, page=None):
    """상단 띠 + 로고 + 제목 + 하단 표시. 본문은 y=1.45in 부터."""
    box(sl, 0, 0, W, Inches(0.12), fill=BLUE); box(sl, 0, Inches(0.12), W, Inches(0.06), fill=GOLD)
    sl.shapes.add_picture(LOGO, Inches(10.9), Inches(0.28), height=Inches(0.75))
    if kicker: box(sl, Inches(0.5), Inches(0.3), Inches(9), Inches(0.35), kicker, 12, False, GRAY)
    box(sl, Inches(0.5), Inches(0.6), Inches(10), Inches(0.8), title, 30, True, BLUE, anchor=MSO_ANCHOR.MIDDLE)
    box(sl, 0, H - Inches(0.1), W, Inches(0.1), fill=BLUE)
    box(sl, Inches(0.5), H - Inches(0.45), Inches(9), Inches(0.3), '대구금향 로타리 클럽 임원 교육 · 신생클럽 어드바이저 심천 김영상', 10, False, GRAY)
    if page: box(sl, W - Inches(1.3), H - Inches(0.45), Inches(0.8), Inches(0.3), str(page), 10, False, GRAY, align=PP_ALIGN.RIGHT)

def notes(sl, text):
    sl.notes_slide.notes_text_frame.text = text

def new(title=None, kicker=None, page=None):
    sl = prs.slides.add_slide(BLANK)
    if title: frame(sl, title, kicker, page)
    return sl

# ─────────────────────────────────────────────────────────────────
# 1 표지
sl = new()
box(sl, 0, 0, W, H, fill=BLUE)
box(sl, 0, Inches(5.9), W, Inches(0.08), fill=GOLD)
sl.shapes.add_picture(LOGO, Inches(0.7), Inches(0.6), height=Inches(1.1))
box(sl, Inches(0.7), Inches(2.0), Inches(11), Inches(0.5), '국제로타리 3700지구 · 2026-27 「지속적인 영향력을」', 16, False, WHITE)
box(sl, Inches(0.7), Inches(2.5), Inches(11.5), Inches(1.3), '클럽 임원의 역할', 60, True, WHITE)
box(sl, Inches(0.7), Inches(3.8), Inches(11.5), Inches(0.8), '대구금향 로타리 클럽 창립 임원 교육', 30, False, WHITE)
box(sl, Inches(0.7), Inches(4.7), Inches(11.5), Inches(0.9), lines=[('“함께하는 리더십이 클럽의 가치를 만듭니다”', 20, False, GOLD)])
box(sl, Inches(0.7), Inches(6.1), Inches(11.5), Inches(0.9), lines=[('강사  신생클럽 어드바이저  심천 김영상  (대구금송로타리클럽 4대 회장 · 2023-24 총재지역대표)', 18, True, WHITE), ('원자료: 3700지구 연수위원회 수경 이명미 「클럽 임원의 역할」(2026-09-16)을 금향 창립 회기에 맞게 재구성', 11, False, RGBColor(0xC8, 0xD3, 0xE8))])
notes(sl, '안녕하십니까. 대구금향 로타리 클럽 신생클럽 어드바이저 심천 김영상입니다. 스폰서클럽인 대구금송로타리클럽 4대 회장을 지냈고, 2023-24년도에는 3700지구 총재지역대표를 맡았습니다. 오늘은 3700지구 연수위원회의 「클럽 임원의 역할」 강의를 우리 금향 창립 회기에 맞게 다시 정리해 말씀드리겠습니다.')

# 2 강사·어드바이저의 역할
sl = new('신생클럽 어드바이저는 무엇을 하는 사람인가', 'SPONSOR CLUB · NEW CLUB ADVISOR', 2)
if os.path.exists(ADVISOR_PHOTO): sl.shapes.add_picture(ADVISOR_PHOTO, Inches(0.6), Inches(1.6), height=Inches(4.2))
box(sl, Inches(4.2), Inches(1.6), Inches(8.5), Inches(1.0), lines=[('심천 김영상', 28, True, BLUE), ('대구금송로타리클럽 4대 회장 · 대구금향RC 신생클럽 어드바이저', 16, False, GRAY)])
box(sl, Inches(4.2), Inches(2.7), Inches(8.5), Inches(4.1), lines=[
    ('이력', 16, True, BLUE),
    ('· 대구금송로타리클럽 4대 회장  · 2023-24년도 3700지구 총재지역대표', 14),
    ('· 로타리재단 고액기부자 레벨 1 · PHS(폴 해리스 소사이어티) 회원', 14),
    ('', 6),
    ('어드바이저가 하는 일', 16, True, BLUE),
    ('① 창립 회기 동안 스폰서클럽(금송)의 운영 경험을 전수합니다.', 15),
    ('② 이사회에 참관해 자문합니다 — 결정은 금향 이사회가 합니다.', 15),
    ('③ 지구·존 행사와 스폰서클럽 사이의 다리가 됩니다.', 15),
    ('④ 초대 회기에는 "직전회장"이 없습니다. 그 빈자리의 경험을 대신 채웁니다.', 15),
    ('', 8),
    ('오늘 강의의 목표', 16, True, BLUE),
    ('직책을 외우는 시간이 아니라, 내 역할을 발견하는 시간', 15),
], spacing=1.3)
notes(sl, '먼저 제 역할부터 말씀드립니다. 어드바이저는 결정하는 사람이 아닙니다. 금송에서 겪은 것을 먼저 이야기해 드리고, 여러분이 결정하실 수 있게 돕는 사람입니다. 신생클럽에는 직전회장이 없습니다. 보통 직전회장이 하는 자문 역할을 창립 회기에는 제가 대신하겠습니다. [확인필요: 스폰서클럽 지원 기간·의무는 RI 신생클럽 규정으로 확인]')

# 3 함께 생각해 볼 질문
sl = new('강의 시작 전, 함께 생각해 볼 세 가지 질문', 'PRE-SESSION', 3)
for i, (q, sub) in enumerate([('초대회장은 무엇을 해야 할까?', '창립 회기, 아직 아무 전통도 없는 클럽에서'),
                              ('총무이사는 무엇을 해야 할까?', '기록도, 회의록도, 명부도 오늘부터 시작'),
                              ('내가 맡은 자리에서', '금향이 더 즐겁고, 활기차고, 좋은 봉사를 하려면 무엇을 해야 할까?')]):
    y = Inches(1.7 + i * 1.75)
    box(sl, Inches(0.6), y, Inches(1.0), Inches(1.0), str(i + 1), 32, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, fill=GOLD, shape=MSO_SHAPE.OVAL)
    box(sl, Inches(1.9), y, Inches(10.5), Inches(1.2), lines=[(q, 24, True, BLUE), (sub, 15, False, GRAY)], anchor=MSO_ANCHOR.MIDDLE)
notes(sl, '회원들은 우리의 직책을 기억할까요? 아닙니다. 우리가 클럽을 위해 무엇을 했는지를 기억합니다. 강의 시작 전 세 가지 질문을 드립니다. 창립 회기라 더 어렵고, 더 중요한 질문입니다.')

# 4 톱니바퀴 — 13석 이사회
sl = new('톱니바퀴의 비밀 — 금향 이사회 13석', 'CHAPTER 01 · THE SYMBOL', 4)
box(sl, Inches(0.6), Inches(1.6), Inches(5.2), Inches(5.0), lines=[
    ('톱니 하나가 빠지면 멈춥니다.', 20, True, BLUE),
    ('각자의 자리에서 역할을 잘하는 사람들이 서로 맞물릴 때 추진력이 생깁니다.', 15),
    ('', 8),
    ('금향 이사회는 13석입니다.', 18, True, BLUE),
    ('회장단 3  ·  직무이사 3  ·  상임위원장 7', 15),
    ('창립 회기에는 초대회장이 차기회장을 겸해 12명이 13석을 채웁니다.', 15),
    ('', 8),
    ('의사결정 원칙', 18, True, BLUE),
    ('회원 의견 수렴 → 이사회 결정 → 전 회원 참여', 15, True),
], spacing=1.3)
cols = [('회장단', ['회장 윤림 이샘결 (초대회장, 차기회장·골프회장 겸)', '부회장 희원 박미애 (문화레저동호회장)']),
        ('직무이사', ['총무이사 채영 전소연 (골프회 총무)', '재무이사 채원 김윤정', '사찰이사 서휘 이은정']),
        ('상임위원장 7', ['클럽관리 화림 이진희', '멤버십 소운 석혜영', '공공이미지 윤재 오재현', '로타리재단 해원 김수영', '봉사프로젝트 혜은 김영미', 'DEI 설화 김은진', 'IT 은수 서은지'])]
y = Inches(1.6)
for name, items in cols:
    h = Inches(0.55 + 0.3 * len(items))
    box(sl, Inches(6.2), y, Inches(6.5), h, lines=[('', 12)] + [(' · ' + t, 12, False, INK) for t in items], fill=None, line=BLUE, spacing=1.2)
    box(sl, Inches(6.2), y, Inches(6.5), Inches(0.38), fill=BLUE)
    box(sl, Inches(6.3), y, Inches(6.3), Inches(0.38), name, 14, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
    y += h + Inches(0.15)
notes(sl, '로타리의 상징 톱니바퀴입니다. 금향의 톱니는 열세 개입니다. 회장단 셋, 직무이사 셋, 위원장 일곱. 국제로타리 권장 5개 위원회에 DEI와 IT를 더해 7개로 확정했고, 이는 금송과 같은 구성입니다. 여러분 한 분 한 분이 톱니입니다. 그리고 결정은 이사회가 합니다. 회원 의견을 모아 이사회가 결정하고, 전 회원이 함께 참여합니다.')

# 5 리더십 승계 — 금향의 특수성
sl = new('리더십 승계 — 신생클럽은 무엇이 다른가', 'LEADERSHIP CONTINUITY', 5)
for i, (t, body) in enumerate([
    ('일반 클럽', '직전회장(자문) → 회장 → 차기회장(준비)\n세 회장이 겹쳐 경험이 이어집니다.'),
    ('금향 창립 회기', '직전회장 없음 · 초대회장이 차기회장 겸임\n초대회장은 1대·2대 회장 역임(창립 회기 + 2027-28)\n다음 차기회장이 3대 회장\n→ 12월 연차총회에서 차차기 회장단을 인준'),
    ('빈자리를 채우는 방법', '스폰서클럽 어드바이저의 자문 · 5월 클럽 리더십 연수회(전 임원·신입회원) · 이사회 기록을 남겨 다음 회장단에 넘기기')]):
    x = Inches(0.6 + i * 4.15)
    box(sl, x, Inches(1.7), Inches(3.9), Inches(0.6), t, 18, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, fill=BLUE if i != 1 else GOLD)
    box(sl, x, Inches(2.3), Inches(3.9), Inches(3.6), body, 15, fill=LIGHT, spacing=1.35)
box(sl, Inches(0.6), Inches(6.1), Inches(12.1), Inches(0.7), '앞선 경험을 이어받아 다음 사람에게 연결하는 것 — 지속 가능한 클럽을 만드는 리더십 승계', 16, True, BLUE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
notes(sl, '보통 클럽은 직전회장·회장·차기회장 세 분이 겹쳐 있어 경험이 자연스럽게 흐릅니다. 금향은 다릅니다. 직전회장이 없고, 초대회장이 차기회장을 겸하며 1대와 2대 회장을 역임합니다. 다음 차기회장이 3대 회장이 됩니다. 그래서 세 가지로 빈자리를 채웁니다. 어드바이저의 자문, 5월 클럽 리더십 연수회, 그리고 기록입니다. 기록은 총무이사와 AI 사무장이 맡습니다.')

# 6 임원 활동의 3가지 기반
sl = new('클럽 임원 활동의 3가지 기반 — 배우고, 투명하게, 연결하고', 'FOUNDATIONS', 6)
for i, (t, items) in enumerate([
    ('교육·연수', ['4월 차기 회장·총무 연수회(PELS)', '5월 클럽 리더십 연수회 — 회장·전 임원·신입회원', '7·9월 지구 세미나(회원증강·공공이미지·재단)', '금향 AI 사무장 교육 봇: 20회차 영상 + 확인질문']),
    ('재정·기부', ['회비 3항목 72만 원: 연회비 30·의무봉사금 30·주회식대 12만', '임원 분담금은 창립 회기 면제(2027-07-01부터)', '지출 3단 분리: 재무 기안 → 회장 승인 → 재무 집행 → 총무 공개', 'RI 기부(RFSM $100)는 2027-28부터 권고 — 창립 회기엔 어드바이저 독려']),
    ('네트워크', ['스폰서클럽 대구금송RC · 신생클럽 어드바이저', '존12 · 총재지역대표 · 지구', '골프회(홀수달) · 문화레저동호회(짝수달)로 회원끼리 연결', '자매클럽·해외 봉사는 다음 단계'])]):
    x = Inches(0.6 + i * 4.15)
    box(sl, x, Inches(1.7), Inches(3.9), Inches(0.6), t, 18, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, fill=BLUE)
    box(sl, x, Inches(2.3), Inches(3.9), Inches(4.2), lines=[('· ' + s, 13) for s in items], fill=LIGHT, spacing=1.4)
notes(sl, '임원 활동의 기반은 세 가지입니다. 배우는 것, 투명하게 관리하는 것, 서로 연결하는 것. 금향은 회비 3항목 72만 원과 지출 3단 분리를 이미 기준으로 정해 두었고, 창립총회 뒤 재무 기능을 가동합니다. 회장 승인 없이 집행할 수 없고, 본인 기안을 본인이 승인할 수 없습니다.')

# 7 임원 3대 핵심 역할
sl = new('클럽 임원의 3대 핵심 역할', 'CORE ROLES', 7)
for i, (t, body) in enumerate([('클럽 운영', '회장·총무·재무·사찰이 맡은 역할을 충실히.\n정기모임·이사회·기록·재정·의전'),
                               ('봉사 활동', '봉사프로젝트위원장과 로타리재단위원장이 대외 활동과 기금 조성을 주도.\n회장 중점사업: 지역 취약계층 봉사 프로젝트 개발(지자체·봉사단체·기관 연계)'),
                               ('회원 관리', '새 회원 영입(목표 25명)과 기존 회원의 참여.\n멤버십·클럽관리·DEI 위원장이 함께')]):
    x = Inches(0.6 + i * 4.15)
    box(sl, x, Inches(1.8), Inches(3.9), Inches(1.0), t, 24, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, fill=GOLD if i == 1 else BLUE)
    box(sl, x, Inches(2.8), Inches(3.9), Inches(3.0), body, 15, fill=LIGHT, spacing=1.4)
box(sl, Inches(0.6), Inches(6.0), Inches(12.1), Inches(0.8), '임원은 직책을 맡은 사람이 아니라, 클럽과 봉사와 회원을 연결하는 사람', 18, True, BLUE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
notes(sl, '크게 보면 임원의 일은 세 가지입니다. 클럽을 운영하고, 봉사하고, 회원을 관리하는 것. 금향의 봉사는 회장 중점사업이 방향입니다. 지역 취약계층을 위한 봉사 프로젝트를 개발하고, 지자체와 봉사단체, 관련 기관과 연계합니다.')

# 역할 슬라이드 공통
def role_slide(page, kicker, title, quote, resp, actions, key, gh):
    sl = new(title, kicker, page)
    box(sl, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.5), quote, 14, False, GRAY)
    box(sl, Inches(0.5), Inches(1.95), Inches(7.6), Inches(0.35), 'KEY RESPONSIBILITIES · 핵심 역할', 11, True, BLUE)
    for i, (t, b) in enumerate(resp):
        y = Inches(2.3 + i * 1.02)
        box(sl, Inches(0.5), y, Inches(0.55), Inches(0.55), '0%d' % (i + 1), 13, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, fill=BLUE)
        box(sl, Inches(1.15), y - Inches(0.05), Inches(6.9), Inches(1.0), lines=[(t, 15, True, INK), (b, 12, False, GRAY)], spacing=1.2)
    box(sl, Inches(8.4), Inches(1.95), Inches(4.4), Inches(0.35), 'ACTION GUIDE · 3가지 행동 원칙', 11, True, BLUE)
    for i, (t, b) in enumerate(actions):
        y = Inches(2.3 + i * 0.95)
        box(sl, Inches(8.4), y, Inches(0.5), Inches(0.5), 'ABC'[i], 13, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, fill=GOLD, shape=MSO_SHAPE.OVAL)
        box(sl, Inches(9.0), y - Inches(0.05), Inches(3.8), Inches(0.95), lines=[(t, 13, True, INK), (b, 11, False, GRAY)], spacing=1.15)
    box(sl, Inches(8.4), Inches(5.2), Inches(4.4), Inches(1.65), lines=[('금향 창립 회기 포인트', 12, True, BLUE)] + [('· ' + g, 11, False, INK) for g in gh], fill=LIGHT, spacing=1.2)
    box(sl, Inches(0.5), Inches(6.45), Inches(7.6), Inches(0.5), '핵심  ' + key, 13, True, BLUE, anchor=MSO_ANCHOR.MIDDLE)
    return sl

sl = role_slide(8, '로타리클럽 임원 역할 · 회장단', '회장의 역할: 방향을 만들고, 사람을 세우는 사람',
    '“회장은 모든 일을 하는 사람이 아닙니다. 회원들이 일하도록 방향을 잡는 리더입니다.”',
    [('클럽 비전 제시', '연간 목표를 세우고 방향과 가치를 전 회원과 공유'), ('회의 주재 및 총괄', '정기모임·이사회를 주재하고 봉사 프로젝트와 주요 활동을 총괄'),
     ('대외적 대표 역할', '지구·지역 회장·총무 협의회, 지역사회 행사에 클럽 대표로 참석'), ('회원 관리 및 성장 지원', '회원의 참여를 독려하고 재능을 발휘하도록 지지·육성')],
    [('명확한 역할 맡기기', '구체적 역할을 부여해 회원 스스로 주인의식을 갖게'), ('이름과 직책으로 인정하기', '공식 석상의 진심 어린 감사와 격려'), ('먼저 도울 것 묻기', '지시 대신 지원 의사를 먼저 물어 자발적 협력을')],
    '좋은 회장은 앞에서 끄는 사람이 아닌, 사람을 세워주는 사람입니다.',
    ['초대회장 1대·2대 회장 역임 · 차기회장·골프회장 겸임', '중점사업: 취약계층 봉사 프로젝트 개발', '12월 연차총회 → 2~3월 지역 모임 → 4월 PELS → 5월 CLR', '회원 의견 수렴 → 이사회 결정 → 전 회원 참여'])
notes(sl, '회장은 모든 일을 하는 사람이 아닙니다. 초대회장은 특히 그렇습니다. 전통이 없으니 모든 것을 혼자 하려다 지치기 쉽습니다. 회원 의견을 모아 이사회가 결정하고 전 회원이 참여하는 구조를 첫 회기에 세우는 것이 초대회장의 가장 큰 일입니다. 부회장은 회장 부재 시 대행하고, 문화레저동호회를 맡습니다.')

sl = role_slide(9, '로타리클럽 임원 역할 · 직무이사', '총무이사의 역할: 소통을 잇고, 클럽을 움직이는 사람',
    '“총무는 행정만 하는 사람이 아닙니다. 클럽이 원활히 돌아가도록 소통과 살림을 책임지는 엔진입니다.”',
    [('정기모임 및 회의 지원', '일정 조율, 회의 준비, 공식 회의록 작성·공람'), ('회원 및 출석 관리', '회원 명부·출석 현황 관리, 신입회원 등록, 지구/RI 데이터 갱신'),
     ('행정 및 대내외 소통', '지구·RI 공지를 회원에게 신속히 전달, 공식 서신'), ('클럽 기록 보존', '회의록·공문·활동 이력을 보관해 운영의 연속성을')],
    [('정확하고 빠른 공유', '일정과 공지를 명확히 전달'), ('세심한 기록과 보존', '작은 회의 결과도 기록 — 공정성과 신뢰'), ('먼저 다가가는 소통', '출석과 안부를 살펴 참여를 이끌기')],
    '좋은 총무는 전달자가 아닌, 클럽을 신뢰로 묶어주는 조력자입니다.',
    ['골프회 총무 겸임(창립 회기)', 'AI 사무장이 보조: 출석·문서 보관(드라이브)·회의 알림', '창립총회 의안·세칙(안)·식순·초청장 담당(준비 탭)', '회의록은 드라이브 03_회의록·총회에'])
notes(sl, '총무이사는 클럽의 엔진입니다. 금향에는 AI 사무장이 있어 출석 기록, 문서 보관, 알림을 대신합니다. 그래서 총무이사는 기록의 정확성과 사람 사이의 소통에 집중하실 수 있습니다. 창립총회 준비 항목 가운데 세칙 초안, 식순, 초청장이 총무이사 몫입니다.')

sl = role_slide(10, '로타리클럽 임원 역할 · 직무이사', '재무이사의 역할: 투명한 관리로 신뢰를 지키는 사람',
    '“재무는 회원의 소중한 회비가 가치 있는 봉사로 이어지게 하는 핵심 조력자입니다.”',
    [('회비 및 기금 관리', '회비 수입·기금·기부금을 정확하게 통합 관리'), ('예산 수립 및 운영', '연간 운영과 봉사사업의 현실적 예산 편성'),
     ('지출 심의 및 관리', '승인된 예산 안에서 지출의 적정성 검토·집행'), ('결산 보고 및 투명성', '정기 재정 보고와 연간 결산으로 신뢰 구축')],
    [('투명한 기록', '모든 수입·지출에 증빙을 갖추고 체계화'), ('예산 공유', '임원진과 협의해 세우고 회원과 공유'), ('정확한 보고', '정기 결산 보고로 건전성을 증명')],
    '투명한 기록이 클럽의 백 번의 신뢰를 만듭니다.',
    ['창립 회기 회비 525,000 = (연회비 30 + 의무봉사금 30 + 주회비 10만 = 70만) × 9/12, 납기 10/14 창립총회일 · 본 회기(27-28)부터 주회비 12만 → 연 72만, 7/1 일시납', '지출 3단 분리 — 본인 기안·본인 승인 금지', '정정은 삭제 없이 정정행 추가(로그)', '개인 미납 안내는 1:1로만 · 월보고 매월 1일 자동'])
notes(sl, '재무이사는 신뢰를 지키는 사람입니다. 금향은 회비 3항목 72만 원과 지출 3단 분리를 기준으로 정했습니다. 재무가 기안하고, 회장이 승인하고, 재무가 집행하고, 총무가 공개합니다. 삭제는 없고 정정행만 남깁니다. 창립총회에서 예산안과 회비안이 확정되면 AI 사무장의 재무 기능이 가동됩니다.')

sl = role_slide(11, '로타리클럽 임원 역할 · 직무이사', '사찰이사의 역할: 규율을 세우고, 회의의 품격을 높이는 사람',
    '“사찰은 규율 관리자가 아닙니다. 품격과 질서를 지키는 파수꾼이자, 친교의 가교입니다.”',
    [('정관 준수 및 회의 질서', '정관·세칙에 따라 회의가 품격 있게 진행되도록 의전 지원'), ('회의 환경 및 의전 점검', '먼저 도착해 좌석·음향·현수막 점검, 정시 출석 유도'),
     ('친교 활성화 및 가교', '위트 있는 사찰 활동과 스피치, 회원 의견을 집행부에 전달'), ('의전 물품', '클럽기·로타리기·타종·배너·명패 준비')],
    [('정확하고 엄정한 기준', '규정에 따라 일관되고 공정하게'), ('세심한 사전 준비', '모임 30분 전 도착해 점검'), ('따뜻하고 유쾌한 친교', '지적에 머물지 않고 위트와 배려로')],
    '좋은 사찰은 클럽을 신뢰와 품격으로 묶어주는 조력자입니다.',
    ['창립총회 의전 물품·리허설(사회·타종·입장 동선) 담당', '첫 회기: 세칙과 함께 금향의 회의 예절을 처음 세우는 자리', '여성 클럽의 품격 — 스폰서클럽 금송 사찰 방식을 참고'])
notes(sl, '사찰이사는 회의의 품격을 지키는 분입니다. 창립총회 리허설과 의전 물품이 사찰이사 몫입니다. 그리고 첫 회기에 세워지는 회의 예절이 금향의 전통이 됩니다. 금송의 사찰 방식을 참고하되 금향답게 만드시면 됩니다.')

# 12~18 위원장 (2열 카드)
def chair_slide(page, title, sub, cards, gh, note):
    sl = new(title, '로타리클럽 임원 역할 · 상임위원장', page)
    box(sl, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.5), sub, 14, False, GRAY)
    for i, (t, b) in enumerate(cards):
        x, y = Inches(0.5 + (i % 2) * 6.2), Inches(2.0 + (i // 2) * 1.75)
        box(sl, x, y, Inches(0.55), Inches(0.55), '0%d' % (i + 1), 13, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, fill=BLUE)
        box(sl, x + Inches(0.65), y - Inches(0.05), Inches(5.4), Inches(1.6), lines=[(t, 15, True, INK), (b, 12, False, GRAY)], spacing=1.2)
    box(sl, Inches(0.5), Inches(5.6), Inches(12.3), Inches(1.3), lines=[('금향 창립 회기 포인트', 12, True, BLUE)] + [('· ' + g, 12, False, INK) for g in gh], fill=LIGHT, spacing=1.2)
    notes(sl, note); return sl

chair_slide(12, '클럽관리위원장 — 클럽의 심장이자 든든한 버팀목', '회원이 즐겁게 참여하고 함께 성장하는 클럽 만들기',
    [('오고 싶고 즐거운 정기모임 운영', '형식적 회의를 넘어 알찬 프로그램과 배움·감동이 있는 소통의 장'), ('회원 친교 및 소통', '월 1회 친교·번개 모임, 회원 사업장 방문, 회원 소식과 봉사 후기 공유'),
     ('단 한 명도 소외되지 않는 관심', '오랜만에 온 회원을 반갑게, 결석한 회원에게 먼저 안부'), ('금향만의 전통 만들기', '회원 감사의 날, 신입회원 환영의 날 — 소속감과 우정')],
    ['정기모임 요일·시간·장소(안)를 창립총회 의안으로', '골프회(홀수달)·문화레저동호회(짝수달)와 나눠 매달 만남을', '주간 점검모임을 정기모임으로 이어 가기'],
    '클럽관리위원장은 정기모임을 오고 싶은 자리로 만드는 분입니다. 금향은 홀수달 골프회, 짝수달 문화레저동호회로 매달 회원이 만납니다. 정기모임 요일·시간·장소는 창립총회 의안입니다.')

chair_slide(13, '멤버십위원장 — 회원을 모으고 지키는 곳', '오고 싶고, 머물고 싶고, 자랑하고 싶은 클럽 만들기',
    [('회원 증강 목표 설정 및 영입', '회원 1인 1명 게스트 초청, 봉사활동에 초청, 로타리 소개 자료 활용'), ('신입회원 교육 및 멘토링', '1:1 멘토 결연, 오리엔테이션, 입회 후 6개월 정착'),
     ('회원 봉사 참여', '전 회원 1인 1봉사 분과 배정, 우수 봉사 회원 격려'), ('회원 밀착 관리 및 탈회 방지', '3회 이상 결석 회원 1:1 소통, 친선 교류 정례화')],
    ['창립회원 목표 25명 — 예비회원 명단은 AI 사무장이 관리(추천 양식 접수)', '회원증강 3대 전략: 영입 대상 다양화 · 기존 회원 유지 · 클럽 매력도 제고', '7월 지구 회원증강·공공이미지 세미나'],
    '멤버십위원장은 창립 클럽에서 가장 중요한 자리 중 하나입니다. 목표는 25명입니다. 예비회원은 추천 양식으로 접수하면 AI 사무장이 명단으로 관리합니다. 3700지구 세미나의 3대 전략을 기준으로 삼으십시오.')

chair_slide(14, '공공이미지위원장 — 클럽을 알리는 곳', '단순 행사 홍보가 아니라, 로타리의 가치와 봉사의 영향력을 지역사회에 알리고 새 회원까지 연결하는 역할',
    [('홍보 활동', '언론 보도, SNS·밴드를 활용해 봉사 활동과 성과를 대외에'), ('지역 사회 소통', '공공기관·사회단체와 파트너십, 긍정적 이미지 제고'),
     ('행사 및 봉사활동 홍보', '주요 행사·사회공헌 프로젝트를 기획·홍보해 참여와 관심을'), ('로타리 브랜드 가치 높이기', "'초아의 봉사' 정신 전파, 명찰·조끼 착용, 로고·색상 규정 준수")],
    ['창립총회: 현수막·포토존·식순지 제작(10/7), 촬영 담당 지정(10/7), 사진 드라이브 보관·공유(10/17)', '클럽 SNS 계정 열기 · 첫 봉사 사업 홍보', '봉사프로젝트위원장·IT위원장과 함께'],
    '공공이미지위원장은 우리 클럽이 하는 좋은 일을 지역사회가 알게 하는 분입니다. 창립총회 홍보물과 촬영이 첫 과제입니다. 로타리 로고와 파란색·금색은 규정대로 씁니다.')

chair_slide(15, '봉사프로젝트위원장 — 봉사를 기획하고 실행하는 곳', '지역사회의 긍정적 변화를 이끌고, 전 회원이 함께 참여해 지속적인 변화를 만드는 역할',
    [('봉사 사업 기획', '지역사회의 필요를 파악하고 클럽 맞춤형 프로젝트 입안(직업봉사·지역사회봉사·청소년)'), ('로타리재단과 연계한 봉사', '지구보조금·글로벌보조금 활용 — 규모와 영향력'),
     ('회원 참여 봉사활동 확대', '연간 봉사계획 사전 안내, 신입회원 첫 봉사 참여, 가족 봉사'), ('성과 평가', '사업 뒤 피드백 수렴, 개선사항 반영, 지속 가능한 봉사로')],
    ['회장 중점사업을 실행: 지역 취약계층 봉사 프로젝트 개발 — 지자체·봉사단체·관련 기관 연계', '7대 초점분야(평화·질병·물위생·모자보건·교육·지역경제·환경) 기준', '봉사 계정은 운영 계정과 분리'],
    '봉사프로젝트위원장은 로타리의 존재 이유를 실제로 만드는 분입니다. 금향의 첫 봉사는 회장 중점사업, 지역 취약계층을 위한 프로젝트입니다. 지자체 한 곳에 묶이지 말고 봉사단체·관련 기관과 넓게 연계하십시오.')

chair_slide(16, '로타리재단위원장 — 기부와 보조금을 맡는 곳', '회원이 재단에 왜 기부하는지 이해하고 참여하도록 돕고, 기부금이 보조금 사업을 통해 봉사로 이어지게 연결하는 역할',
    [('기금 조성 — 전 회원 연차기금 참여', 'Every Rotarian, Every Year · PHF·RFSM 안내 · 전 회원 재단교육 연 2회'), ('보조금 사업 발굴·추진', '지구보조금·글로벌보조금 신청·승인 절차, 진행 모니터링'),
     ('지역사회 봉사와의 연계', '클럽 봉사와 재단 7대 초점분야 사업을 유기적으로 연계'), ('투명한 기금 관리', '보조금·재단 기금 사용 내역 공개, 영수증 증빙, 결산 보고')],
    ['RI 기부(RFSM US$100)는 2027-28부터 권고 — 창립 회기 청구 없음', '창립 회기 미션: 어드바이저와 함께 기부 방식을 안내·독려 · PHF(미화 1,000달러)', '지구보조금 자격 요건을 지구에 확인 [확인필요]'],
    '재단위원장은 기부와 보조금을 맡습니다. RI 기부 100달러는 2027-28 회기부터 권고 사항으로 두었습니다. 창립 회기에는 제가 이사회에서 기부 방식을 안내하고 독려하는 것을 미션으로 삼겠습니다. 첫 회기 목표는 회원 전원이 재단이 무엇인지 아는 것입니다. 신생클럽의 지구보조금 자격은 지구에 확인이 필요합니다.')

chair_slide(17, 'DEI위원장 — 다양성·공평·포용을 지키는 곳', '누구나 환영받고, 존중받고, 참여할 수 있는 클럽 문화를 만드는 역할 (금향 6번째 위원회)',
    [('환영하는 문화', '신입회원·게스트가 첫날부터 소속감을 느끼는 모임 만들기'), ('공평한 참여', '연령·직업·경력에 관계없이 역할과 발언 기회를 고르게'),
     ('포용의 언어', '회의·단톡·홍보물에서 배제하는 표현 없애기'), ('다양성이 곧 매력', '영입 대상 다양화(멤버십)와 손잡기')],
    ['여성 클럽으로서 금향의 정체성과 다양성을 함께 — 세대·직업의 폭을 넓히기', 'RI DEI 선언을 첫 회기 교육 회차에 포함', '멤버십·클럽관리위원장과 한 팀'],
    '금향은 국제로타리 권장 5개 위원회에 DEI와 IT를 더했습니다. DEI위원장은 누구나 환영받는 문화를 만드는 분입니다. 여성 클럽이라는 정체성 안에서 세대와 직업의 다양성을 넓히는 것이 과제입니다.')

chair_slide(18, 'IT위원장 — 디지털로 클럽을 잇는 곳', '클럽의 소통·기록·홍보 도구를 관리해 임원과 회원의 시간을 아끼는 역할 (금향 7번째 위원회)',
    [('AI 사무장(텔레그램 봇) 운영', '회장단방·임원방 운영, 권한 등록, 출석·교육·준비 체크리스트'), ('공유드라이브 관리', '폴더 체계(창립·정관 / 회원명부 / 회의록 / 재무 / 봉사 / 행사사진 …), 권한'),
     ('온라인 채널', '클럽 SNS·밴드 계정, 회원 단톡 — 공공이미지위원장과 함께'), ('데이터 보호', '회원 연락처·생년은 시트에만, 코드·단톡에 평문 금지')],
    ['AI 사무장 관리위원(동우 김종만)에게서 운영을 이어받는 것이 첫 회기 목표', '3700지구도 클럽 IT위원회 신설을 권장', '지구·RI 시스템(My Rotary) 회원 등록 지원'],
    'IT위원장은 금향의 도구를 맡는 분입니다. AI 사무장, 공유드라이브, 온라인 채널입니다. 지금은 관리위원이 운영하지만 첫 회기 안에 IT위원장이 이어받는 것이 목표입니다. 회원 개인정보는 시트에만 두고 단톡이나 코드에 평문으로 올리지 않습니다.')

# 19 창립 회기 일정
sl = new('금향 창립 회기 — 임원이 함께 걷는 일정', '2026-27 TIMELINE', 19)
steps = [('10/14', '창립총회', '지구회관 5층 19:00\n세칙·임원·예산·회비·정기모임 의결'), ('11~12월', '첫 봉사·정기모임', '회장 중점사업 착수\n동호회 첫 모임'),
         ('12월', '연차총회', '차차기 회장단·이사진 인준'), ('2~3월', '지역 회장·총무\n모임', '총재지역대표 주관'), ('4월', 'PELS·지구대회', '차기 회장·총무 연수회'),
         ('5월', '클럽 리더십\n연수회', '회장·전 임원·신입회원'), ('6/30', '창립 회기 종료', '기록·재정 정리 → 2027-28 회기')]
n = len(steps); x0, x1 = Inches(0.9), Inches(12.4); ybar = Inches(3.0)
sl.shapes.add_connector(1, x0, ybar, x1, ybar).line.color.rgb = GOLD
for i, (m, t, b) in enumerate(steps):
    x = x0 + int((x1 - x0) * i / (n - 1))
    box(sl, x - Inches(0.22), ybar - Inches(0.22), Inches(0.44), Inches(0.44), fill=GOLD if i == 0 else WHITE, line=GOLD, shape=MSO_SHAPE.OVAL)
    box(sl, x - Inches(0.9), ybar - Inches(0.9), Inches(1.8), Inches(0.5), m, 14, True, BLUE, PP_ALIGN.CENTER, MSO_ANCHOR.BOTTOM)
    box(sl, x - Inches(0.9), ybar + Inches(0.35), Inches(1.8), Inches(2.0), lines=[(t, 13, True, INK), (b, 10, False, GRAY)], align=PP_ALIGN.CENTER, spacing=1.15)
box(sl, Inches(0.6), Inches(5.9), Inches(12.1), Inches(0.9), '창립총회 준비 24항목은 AI 사무장 「준비」 탭에서 /prep 으로 확인 — 담당·기한·완료 표시', 14, False, GRAY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, fill=LIGHT)
notes(sl, '창립 회기 일정입니다. 10월 14일 창립총회에서 세칙, 임원, 예산, 회비, 정기모임을 의결합니다. 12월 연차총회에서는 차차기 회장단을 인준합니다. 4월과 5월의 연수회는 임원 전원이 받습니다. 준비 항목은 AI 사무장이 담당과 기한을 매일 알립니다.')

# 20 마무리 — 빈 의자
sl = new('우리의 역할은 다르지만, 목표는 하나입니다', 'SESSION SUMMARY', 20)
box(sl, Inches(0.6), Inches(1.6), Inches(6.0), Inches(2.4), lines=[('작은 관심이 빈 의자를 채웁니다', 20, True, BLUE), ('“요즘 안 보이셔서 궁금했습니다. 다음 모임에는 꼭 얼굴 뵈어요.”', 15, False, INK), ('빈 의자가 채워집니다 — 회원의 참여가 곧 클럽의 활력입니다.', 13, False, GRAY)], fill=LIGHT, spacing=1.4)
for i, (t, b) in enumerate([('회장', '사람을 세움'), ('총무', '사람을 연결'), ('재무', '신뢰 구축'), ('사찰', '올바름 유지'), ('위원장', '동력 제공')]):
    x = Inches(6.9 + i * 1.2)
    box(sl, x, Inches(1.6), Inches(1.1), Inches(1.1), t, 14, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, fill=BLUE, shape=MSO_SHAPE.OVAL)
    box(sl, x - Inches(0.1), Inches(2.75), Inches(1.3), Inches(0.6), b, 11, False, GRAY, PP_ALIGN.CENTER)
box(sl, Inches(0.6), Inches(4.3), Inches(12.1), Inches(2.4), lines=[('실천 가이드', 14, True, BLUE),
    ('✓ 부재 중인 회원에게 먼저 전화하기      ✓ 신입 회원에게 따뜻하게 먼저 다가가기', 14), ('✓ 봉사한 회원에게 진심 어린 감사 전하기   ✓ 모두가 함께 참여하는 긍정적 분위기 조성', 14),
    ('', 6), ('내일부터 실천할 3가지: 내 역할의 매뉴얼(임원 직책 안내) 숙지 · 월 1회 다른 임원과 소통 · 회원이 체감할 작은 변화 하나', 13, True, INK)], spacing=1.4)
notes(sl, '임원의 역할은 결국 사람입니다. 회장은 사람을 세우고, 총무는 연결하고, 재무는 신뢰를 쌓고, 사찰은 올바름을 지키고, 위원장은 동력을 줍니다. 내일부터 세 가지만 하십시오. 임원 직책 안내 문서를 읽고, 다른 임원과 소통하고, 작은 변화를 하나 만드십시오.')

# 21 끝
sl = new()
box(sl, 0, 0, W, H, fill=BLUE); box(sl, 0, Inches(5.2), W, Inches(0.08), fill=GOLD)
box(sl, Inches(0.7), Inches(1.6), Inches(12), Inches(1.0), '임원의 진정한 역할: 함께 걷는 길', 40, True, WHITE)
box(sl, Inches(0.7), Inches(2.8), Inches(12), Inches(2.2), lines=[('처음 문을 열고 들어왔을 때, 낯설어하던 나에게 누군가 건넨 따뜻한 인사 한마디를 기억하시나요?', 18, False, WHITE),
    ('임원의 진정한 역할은 새로 온 회원이 소외감을 느끼지 않도록 곁을 내어주는 것입니다.', 18, False, WHITE), ('“이곳에 오길 참 잘했다는 마음을 전하는 일”', 22, True, GOLD)], spacing=1.5)
box(sl, Inches(0.7), Inches(5.5), Inches(12), Inches(1.2), lines=[('내 역할을 알고, 서로 소통하고, 작은 것부터 행동하자', 20, True, WHITE), ('대구금향 로타리 클럽 · 신생클럽 어드바이저 심천 김영상 · 스폰서클럽 대구금송로타리클럽', 12, False, RGBColor(0xC8, 0xD3, 0xE8))])
notes(sl, '(부드럽고 따뜻한 어조로) 여러분, 임원이라는 자리가 무겁게 느껴질 수 있습니다. 하지만 창립 회원 모두가 처음 이곳에 온 사람들입니다. 서로가 서로에게 따뜻한 인사 한마디를 건네는 것, 그것이 금향의 첫 전통이 되기를 바랍니다. 금송이 함께 걷겠습니다.')

os.makedirs(OUT_DIR, exist_ok=True)
prs.save(OUT); print('saved', OUT, len(prs.slides), 'slides')
